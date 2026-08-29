import openpyxl
import pandas as pd
import pytest

from exporter import (build_average_summary, build_summary_row,
                      export_results, export_timeline)


def make_record(cycle, susceptible=0, infected=0, grave=0, recovered=0,
                dead=0, hospitalized=0, waiting=0, capacity=2):
    """Arma un registro de `history` igual al que produce `_record_history`."""
    return {
        "cycle": cycle, "susceptible": susceptible, "infected": infected,
        "grave": grave, "recovered": recovered, "dead": dead,
        "hospitalized": hospitalized, "waiting": waiting,
        "free_beds": capacity - hospitalized, "capacity": capacity,
        "saturated": capacity - hospitalized <= 0,
    }


def make_history():
    """Historia corta de 3 ciclos: el hospital se satura en el ciclo 1."""
    return [
        make_record(0, susceptible=9, infected=1),
        make_record(1, susceptible=7, infected=1, grave=2, hospitalized=2,
                    waiting=1),
        make_record(2, susceptible=7, recovered=2, dead=1, hospitalized=0),
    ]


def make_params(beds=2):
    return {"escenario": f"{beds} camas", "seed": 42, "initial_beds": beds,
            "population": 10, "initial_infected": 1,
            "transmission_probability": 0.2, "infection_radius": 30,
            "pct_grave": 0.3, "infection_duration": 5, "grave_duration": 4,
            "mortality_hospitalized": 0.1, "mortality_waiting": 0.5}


# ----------------------------------------------------------------------
# export_timeline
# ----------------------------------------------------------------------

def test_export_timeline_writes_sheet_with_spanish_columns(tmp_path):
    path = tmp_path / "salida.xlsx"

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df = export_timeline(make_history(), "timeline_run_1", writer)

    assert list(df.columns) == [
        "ciclo", "susceptibles", "infectados", "graves", "hospitalizados",
        "en_espera", "recuperados", "fallecidos", "camas_ocupadas",
        "camas_libres", "capacidad", "saturado",
    ]
    assert "timeline_run_1" in openpyxl.load_workbook(path).sheetnames


def test_export_timeline_derives_occupied_beds_from_capacity(tmp_path):
    with pd.ExcelWriter(tmp_path / "salida.xlsx", engine="openpyxl") as writer:
        df = export_timeline(make_history(), "timeline_run_1", writer)

    assert list(df["camas_ocupadas"]) == [0, 2, 0]
    assert list(df["camas_libres"]) == [2, 0, 2]


def test_export_timeline_keeps_one_row_per_cycle(tmp_path):
    history = make_history()

    with pd.ExcelWriter(tmp_path / "salida.xlsx", engine="openpyxl") as writer:
        df = export_timeline(history, "timeline_run_1", writer)

    assert len(df) == len(history)
    assert list(df["ciclo"]) == [0, 1, 2]


def test_export_timeline_with_empty_history_writes_only_headers(tmp_path):
    with pd.ExcelWriter(tmp_path / "salida.xlsx", engine="openpyxl") as writer:
        df = export_timeline([], "timeline_run_1", writer)

    assert df.empty
    assert "ciclo" in df.columns


# ----------------------------------------------------------------------
# build_summary_row
# ----------------------------------------------------------------------

def test_summary_row_renames_parameters_to_spanish():
    row = build_summary_row(make_params(beds=5), make_history())

    assert row["camas"] == 5
    assert row["semilla"] == 42
    assert row["poblacion"] == 10
    assert row["mortalidad_sin_cama"] == 0.5


def test_summary_row_reports_final_outcome_counts():
    row = build_summary_row(make_params(), make_history())

    assert row["ciclos_simulados"] == 2
    assert row["fallecidos"] == 1
    assert row["recuperados"] == 2
    assert row["nunca_contagiados"] == 7
    assert row["contagiados_totales"] == 3


def test_summary_row_computes_peaks():
    row = build_summary_row(make_params(), make_history())

    assert row["pico_infectados"] == 1
    assert row["pico_graves"] == 2
    assert row["pico_camas_ocupadas"] == 2
    assert row["pico_en_espera"] == 1


def test_summary_row_measures_saturation_and_unattended_load():
    row = build_summary_row(make_params(), make_history())

    # Solo el ciclo 1 tiene las 2 camas ocupadas y 1 persona esperando.
    assert row["ciclos_saturado"] == 1
    assert row["ciclos_con_pacientes_sin_cama"] == 1
    assert row["espera_acumulada_persona_ciclo"] == 1
    assert row["pct_tiempo_saturado"] == round(1 / 3, 4)


def test_summary_row_computes_mortality_rates():
    row = build_summary_row(make_params(), make_history())

    assert row["tasa_mortalidad_poblacion"] == round(1 / 10, 4)
    assert row["tasa_letalidad_contagiados"] == round(1 / 3, 4)


def test_summary_row_with_empty_history_keeps_parameters_only():
    row = build_summary_row(make_params(beds=7), [])

    assert row["camas"] == 7
    assert "fallecidos" not in row


def test_summary_row_ignores_non_scalar_extra_parameters():
    params = make_params()
    params["random_fn"] = lambda: 0.5

    row = build_summary_row(params, make_history())

    assert "random_fn" not in row


# ----------------------------------------------------------------------
# export_results
# ----------------------------------------------------------------------

@pytest.fixture
def workbook(tmp_path):
    runs = [(make_params(beds=b), make_history()) for b in (5, 10, 20)]
    path = tmp_path / "comparacion.xlsx"
    export_results(runs, path)
    return path


def test_export_results_creates_summary_sheet_first(workbook):
    assert openpyxl.load_workbook(workbook).sheetnames == [
        "resumen", "timeline_run_1", "timeline_run_2", "timeline_run_3",
    ]


def test_export_results_writes_one_summary_row_per_scenario(workbook):
    resumen = pd.read_excel(workbook, sheet_name="resumen")

    assert len(resumen) == 3
    assert list(resumen["camas"]) == [5, 10, 20]
    assert list(resumen["escenario"]) == ["5 camas", "10 camas", "20 camas"]


def test_export_results_embeds_charts_for_the_presentation(workbook):
    book = openpyxl.load_workbook(workbook)

    assert len(book["resumen"]._charts) == 2
    assert len(book["timeline_run_1"]._charts) == 1


def test_export_results_returns_the_output_path(tmp_path):
    path = tmp_path / "comparacion.xlsx"

    assert export_results([(make_params(), make_history())], path) == path


# ----------------------------------------------------------------------
# Promedio de repeticiones
# ----------------------------------------------------------------------

def test_average_summary_groups_repetitions_by_bed_count():
    runs = [(make_params(beds=b), make_history())
            for b in (5, 5, 10)]
    summary = pd.DataFrame([build_summary_row(p, h) for p, h in runs])

    promedios = build_average_summary(summary)

    assert list(promedios["camas"]) == [5, 10]
    assert list(promedios["repeticiones"]) == [2, 1]


def test_average_summary_averages_the_numeric_results():
    history_sin_muertos = [make_record(0, susceptible=10),
                           make_record(1, susceptible=10)]
    summary = pd.DataFrame([
        build_summary_row(make_params(beds=5), make_history()),
        build_summary_row(make_params(beds=5), history_sin_muertos),
    ])

    promedios = build_average_summary(summary)

    # Una corrida con 1 fallecido y otra con 0 -> promedio 0.5
    assert promedios.loc[0, "fallecidos"] == 0.5


def test_export_results_adds_averages_sheet_when_requested(tmp_path):
    runs = [(make_params(beds=b), make_history()) for b in (5, 5, 10, 10)]
    path = tmp_path / "comparacion.xlsx"

    export_results(runs, path, promediar=True)

    book = openpyxl.load_workbook(path)
    assert book.sheetnames[:2] == ["resumen", "promedios"]
    # Los gráficos comparativos se mueven a `promedios`, que es la hoja
    # que hay que mirar cuando hay repeticiones.
    assert len(book["promedios"]._charts) == 2
    assert len(book["resumen"]._charts) == 0


def test_export_results_without_averages_keeps_charts_in_summary(workbook):
    book = openpyxl.load_workbook(workbook)

    assert "promedios" not in book.sheetnames
    assert len(book["resumen"]._charts) == 2
